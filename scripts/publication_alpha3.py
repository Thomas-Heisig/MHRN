"""Build and verify complete edition 1.5 without promoting experiment DATA."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.empirical_campaign import verified_result  # noqa: E402
from scripts.publication_alpha3_text import (  # noqa: E402
    ARCHITECTURE,
    DISCUSSION,
    METHODS,
    REFERENCES,
    REVIEW,
    decode_text,
)

EDITION = "2026-09-13_recursive-epistemics_v1.5"
PREVIOUS = "2026-09-10_recursive-epistemics_v1.4"
TITLE = "Rekursive Epistemik in verkoerperten spikenden neuronalen Architekturen"


def load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def save(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def files(path: Path) -> dict[str, str]:
    return {
        str(p.relative_to(path)): sha(p)
        for p in sorted(path.rglob("*"))
        if p.is_file() and p.name != "manifest.json"
    }


def value(number: float) -> str:
    return f"{number:.7g}"


def assessment(name: str, kind: str) -> str:
    if kind == "boundary_audit":
        return "Grenz-/Vertragsaudit; kein direkter experimenteller Test der inhaltlichen Hypothese."
    if "brian2" in name:
        return "Strenger diskreter Modellvergleich. Ein verfehltes Konformitaetskriterium bleibt ein negativer Befund, kein Laufzeitfehler."
    if "scaling" in name or "performance" in name:
        return "Gemessene technische Last im angegebenen Fenster; keine Extrapolation auf Vollplastizitaet, biologische Netze, GPU oder exklusive Hardware. Die beiden active_scaling-Labels rufen im aktuellen Code dieselbe korrigierte Funktion auf."
    if "dimens" in name or "topology" in name:
        return "Aufgaben- und graphgebundener Vergleich; Geometrie, Motive und reine Labels unterscheiden. Kein universeller Dimensionsvorteil."
    if "independent_replication" in name:
        return "Reproduktionsprotokoll im selben Projekt; der Name belegt kein unabhaengiges Team."
    if "interference" in name:
        return "Begrenzter Interferenz-Screen; getrennte Task-Netze beweisen keine Resistenz eines gemeinsam trainierten Netzes gegen Vergessen."
    if any(token in name for token in ("memory", "world_model", "behavior_profile")):
        return "Kontrollierte technische Komponentenfunktion; keine automatische Attribution an synaptisches Lernen oder subjektive Kognition."
    if "msba" in name:
        return "Experimentelle Gateway-/Ressourcenfunktion. Proxy-Energie ist nicht gemessene physikalische Energie; keine produktive Tool-Use-Freigabe."
    if any(token in name for token in ("embodied", "connectome")):
        return "Synthetischer sensorimotorischer Screen; externe Controllerbeitraege und Beobachtungsgrenzen beachten. Keine biologische Vollsimulation."
    if "association" in name:
        return "Native angeleitete Assoziation mit gehaltenen Testepisoden; Reichweite auf Aufgabe, Seed-Einheit und Kontrollen begrenzt."
    return "Begrenzte protokollgebundene Beobachtung; Primarmessung, Kontrollen und technische Ausfuehrung nicht mit akzeptierter Evidenz gleichsetzen."


def results_text(plan: dict[str, Any], summary: dict[str, Any]) -> str:
    counts = Counter(s.get("execution_kind", "unspecified") for s in plan["selections"])
    lines = [
        "# Ergebnisse der aktuellen Gesamtkampagne",
        "",
        f"Kampagne: **{summary['campaign']}**. Quellcommit: `{summary['source']}`. Quelldigest: `{summary['source_digest']}`.",
        "",
        f"Ausfuehrungsstatus: `{summary['status_counts']}`; **{summary['total_runs']} gespeicherte Seed-/Bedingungsdatensaetze**. Dies ist keine Anzahl unabhaengig bestaetigter Hypothesen.",
        "",
        f"Kategorien: `{dict(counts)}`. Menschliche Vorlagen: **{len(plan['human_templates'])}**, nicht automatisch bearbeitet. Akzeptierte EVID aus dieser Kampagne: **keine**.",
        "",
        "## Gepaarte Kontraste",
        "",
        "Die folgenden Werte werden aus den gespeicherten Analysen uebernommen. Intervalle sind punktweise Bootstrap-Perzentilintervalle; die Holm-Korrektur gilt innerhalb der jeweiligen Kontrastfamilie, nicht fuer die gesamte explorative Kampagne.",
        "",
    ]
    for family, contrasts in summary.get("comparisons", {}).items():
        lines += [f"### {family}", ""]
        for item in contrasts:
            lines += [
                f"**{item['reference']} minus {item['control']}**: n = {item['n_seed_pairs']} gepaarte Seeds; mittlere Differenz {value(item['mean_difference'])}; 95%-Intervall {item['bootstrap_percentile_95_ci']}; p exakt {value(item['exact_two_sided_sign_flip_p'])}; p Holm {value(item['holm_adjusted_p'])}.",
                "",
            ]
    lines += [
        "## Kernaussage und Grenzen",
        "",
        "Die native Assoziationsaufgabe liefert unter den erfassten Kontrollen einen positiven aufgabenspezifischen Lernbefund. Die Dimensionskontraste tragen keinen belastbaren allgemeinen Vorteil von 5D. Der strenge Brian2-Abgleich verfehlt das deklarierte Konformitaetskriterium. Diese Befunde sind voneinander getrennt zu bewerten; keiner etabliert Bewusstsein, biologische Gleichwertigkeit oder allgemeine Intelligenz.",
        "",
        "## Vollstaendige Protokollbilanz",
        "",
        "n bezeichnet gespeicherte Zeilen pro Bedingung. Primarmaesse werden als Mittel [Minimum; Maximum], boolesche Ergebnisse als true/n ausgegeben. Strukturierte oder nicht vorhandene Endpunkte werden nicht erfunden. Vollstaendige Rohdaten und Receipts sind Teil der Kampagne.",
        "",
    ]
    by_name = {item["protocol"]: item for item in summary["protocols"]}
    for index, spec in enumerate(plan["selections"], 1):
        item = by_name[spec["protocol"]]
        kind = spec.get("execution_kind", "unspecified")
        lines += [
            f"### {index:02d}. {item['protocol']}",
            "",
            f"Status: **{item['status']}**; Kategorie: `{kind}`; gespeicherte Laeufe: {item['runs']}; Seeds: `{spec['seeds']}`; Ticks: `{spec.get('ticks')}`.",
            "",
            assessment(item["protocol"], kind),
            "",
        ]
        if item.get("error"):
            lines += [
                "Fehlerbericht: " + str(item["error"]).replace("\n", " ")[-1800:],
                "",
            ]
        for condition, group in item.get("conditions", {}).items():
            entries = []
            for metric in spec.get("primary_outcomes", []):
                number = group.get("metrics", {}).get(metric)
                boolean = group.get("boolean_outcomes", {}).get(metric)
                if number:
                    entries.append(
                        f"{metric}: {value(number['mean'])} [{value(number['min'])}; {value(number['max'])}]"
                    )
                elif boolean is not None:
                    entries.append(
                        f"{metric}: {sum(v is True for v in boolean)}/{len(boolean)} true"
                    )
                else:
                    entries.append(
                        f"{metric}: strukturiert oder nicht ausgewiesen; siehe Rohdaten"
                    )
            lines += [
                f"**{condition}** (n = {group['runs']}): "
                + "; ".join(
                    entries
                    or ["Strukturierte Basisfamilien; siehe Rohdaten und summary.json."]
                ),
                "",
            ]
        lines += [
            f"Artefaktpfad: `research/experiments/{summary['campaign']}/{index:03d}-{spec['protocol']}/`.",
            "",
        ]
    return "\n".join(lines)


def coverage_text(plan: dict[str, Any]) -> str:
    lines = [
        "# Abdeckung, offene Voraussetzungen und Reproduktion",
        "",
        "## Nicht durch automatische Ausfuehrung geschlossene Fragen",
        "",
        "Die folgende Liste unterscheidet fehlende unmittelbare Messprotokolle von fehlenden Registry-Eintraegen. Auch ein vorhandener Grenzvertrag ist kein direktes Experiment.",
        "",
    ]
    for name in plan.get("questions_without_specific_runnable_protocol", []):
        lines.append(
            f"- `{name}`: spezifische Messung oder menschliche Entscheidung bleibt erforderlich."
        )
    lines += ["", "## Menschliche Vorlagen", ""]
    for template in plan["human_templates"]:
        lines.append(
            f"- `{template.get('protocol', template.get('id', 'template'))}`: {template.get('research_question', '')}; nicht automatisch ausgefuellt."
        )
    lines += [
        "",
        "## Reproduktion",
        "",
        "Den exakten Quellcommit aus plan.json in einem sauberen separaten Checkout verwenden. Die aufgezeichneten Paketversionen und Ressourcenbedingungen pruefen. Ein neuer Ausgabepfad ist zwingend; historische DATA duerfen nicht ueberschrieben werden.",
        "",
        "```bash",
        "python -m pip install -e '.[dev,docs,empirical]'",
        "python scripts/empirical_campaign.py --output /absolute/new/run --campaign-id NEW-ID",
        "python scripts/empirical_campaign.py --output /absolute/new/run --analyze",
        "python scripts/run_stage3_reference.py --output /absolute/new/stage3.json",
        "python scripts/publication_alpha3.py --verify",
        "```",
        "",
        "Eine Wiederholung mit geaenderten Quellen ist eine neue Ausfuehrung. Paketversionen und Simulationsergebnisse werden nicht rueckwirkend dem neuen Commit zugeschrieben.",
        "",
    ]
    return "\n".join(lines)


def style_docx(path: Path, title: str, chapter_titles: list[str]) -> None:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document(path)
    for style_name in ("Header", "Footer"):
        if style_name not in doc.styles:
            doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin, section.bottom_margin = Inches(0.85), Inches(0.85)
        section.left_margin, section.right_margin = Inches(0.95), Inches(0.95)
        section.header_distance, section.footer_distance = Inches(0.35), Inches(0.35)
        hp = section.header.paragraphs[0]
        hp.text = "MHRN | Rekursive Epistemik | Fassung 1.5"
        hp.style = doc.styles["Header"]
        fp = section.footer.paragraphs[0]
        fp.alignment = 2
        fp.add_run("Thomas Heisig | 13.09.2026 | ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        fp._p.append(field)
    doc.paragraphs[0].style = doc.styles["Title"]
    doc.styles["Title"].font.name = "Liberation Sans"
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Title"].paragraph_format.space_before = Pt(20)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Liberation Serif", Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name in ("Body Text", "First Paragraph", "Compact"):
        if name in doc.styles:
            doc.styles[name].font.name = "Liberation Serif"
            doc.styles[name].font.size = Pt(11)
            doc.styles[name].paragraph_format.space_after = Pt(6)
            doc.styles[name].paragraph_format.line_spacing = 1.12
    for level, size in ((1, 17), (2, 13), (3, 11.5), (4, 11)):
        style = next(item for item in doc.styles if item.style_id == f"Heading{level}")
        style.font.name, style.font.size = "Liberation Sans", Pt(size)
        style.font.color.rgb = RGBColor.from_string("18384C")
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = False
    for name in ("Header", "Footer"):
        doc.styles[name].font.size = Pt(9)
        doc.styles[name].font.name = "Liberation Sans"
    for name in ("Source Code", "Verbatim Char"):
        if name in doc.styles:
            doc.styles[name].font.name = "Liberation Mono"
            doc.styles[name].font.size = Pt(9)
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        n = len(table.columns)
        widths = [6.6 / n] * n
        if n == 2:
            widths = [2.1, 4.5]
        elif n == 3:
            widths = [1.65, 1.6, 3.35]
        for row_no, row in enumerate(table.rows):
            props = row._tr.get_or_add_trPr()
            if row_no == 0:
                repeat = OxmlElement("w:tblHeader")
                props.append(repeat)
            for ci, cell in enumerate(row.cells):
                cell.width = Inches(widths[min(ci, n - 1)])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + side)
                    margin.set(qn("w:w"), "75")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                tcpr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    for run in p.runs:
                        run.font.size = Pt(10)
                        if row_no == 0:
                            run.bold = True
                        run.text = run.text.replace("_", "_\u200b").replace(
                            "/", "/\u200b"
                        )

    # A populated, bookmark-linked chapter index replaces an empty TOC field.
    # Both the index and the real chapter paragraphs use stable bookmark names.
    def plain(text: str) -> str:
        return " ".join(
            text.replace("\u00a0", " ").replace("*", "").replace("`", "").split()
        )

    paragraphs = list(doc.paragraphs)
    chapter_paragraphs = []
    position = 1
    for chapter_title in chapter_titles:
        found = next(
            (
                i
                for i in range(position, len(paragraphs))
                if paragraphs[i].style.style_id == "Heading1"
                and plain(paragraphs[i].text) == plain(chapter_title)
            ),
            None,
        )
        if found is None:
            raise ValueError("Missing printed chapter: " + chapter_title)
        paragraph = paragraphs[found]
        chapter_paragraphs.append(paragraph)
        paragraph.paragraph_format.page_break_before = True
        position = found + 1
    anchor = chapter_paragraphs[0]
    toc_title = anchor.insert_paragraph_before("Inhaltsverzeichnis")
    toc_title.paragraph_format.page_break_before = True
    toc_title.paragraph_format.keep_with_next = True
    for run in toc_title.runs:
        run.font.name, run.font.size, run.bold = "Liberation Sans", Pt(17), True
    for index, (chapter_title, paragraph) in enumerate(
        zip(chapter_titles, chapter_paragraphs), 1
    ):
        bookmark = "mhrn_chapter_" + str(index)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(20000 + index))
        start.set(qn("w:name"), bookmark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(20000 + index))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)
        entry = anchor.insert_paragraph_before()
        entry.paragraph_format.space_after = Pt(5)
        entry.paragraph_format.keep_together = True
        entry.paragraph_format.line_spacing = 1.05
        entry.paragraph_format.tab_stops.add_tab_stop(
            Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), bookmark)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = f"{index:02d}. {chapter_title}"
        run.append(text)
        link.append(run)
        entry._p.append(link)
        entry.add_run("\t")
        page_ref = OxmlElement("w:fldSimple")
        page_ref.set(qn("w:instr"), "PAGEREF " + bookmark + " \\h")
        entry._p.append(page_ref)
    # Current and historical part labels stay with the next real chapter.
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(("Teil I -", "Teil II -")):
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
    # Avoid title-only separator pages: the first chapter follows its part label.
    if len(chapter_paragraphs) > 8:
        chapter_paragraphs[8].paragraph_format.page_break_before = False
    # Prevent code paths and digests from crossing the right text margin.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if len(run.text) > 28 and " " not in run.text.strip():
                run.text = re.sub(r"([/_-])", "\\1\u200b", run.text)
                run.text = re.sub(
                    r"([a-f0-9]{20})(?=[a-f0-9]{20})", "\\1\u200b", run.text
                )
    for shape in doc.inline_shapes:
        if shape.width > Inches(6.6):
            shape.height = int(shape.height * Inches(6.6) / shape.width)
            shape.width = Inches(6.6)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    doc.core_properties.title = title
    doc.core_properties.author = "Thomas Heisig"
    doc.core_properties.subject = (
        "Source-bound scientific manuscript; human review pending"
    )
    doc.save(path)


def export_docx(folder: Path) -> None:
    for source, target, title in (
        (
            "MANUSCRIPT.md",
            "MHRN_Dissertationsmanuskript_v1.5.docx",
            "Rekursive Epistemik - Dissertationsmanuskript 1.5",
        ),
        (
            "FORSCHUNGSBERICHT.md",
            "MHRN_Forschungsbericht_v1.5.docx",
            "MHRN - Forschungsbericht 1.5",
        ),
    ):
        text = (folder / source).read_text(encoding="utf-8")
        text = re.sub(r"^\[Inhalts[^\n]+\n", "", text, flags=re.M)
        text = re.sub(r'<a\s+id="[^\"]*"></a>\s*\n', "\n\n", text)
        text = text.replace("\n# Teil I - Aktueller wissenschaftlicher Stand\n\n", "\n")
        printable = folder / ("_print_" + source)
        printable.write_text(text, encoding="utf-8")
        subprocess.run(
            [
                "pandoc",
                str(printable),
                "--resource-path",
                str(folder) + ":" + str(ROOT / "research/publications" / PREVIOUS),
                "--from",
                "markdown+tex_math_dollars",
                "--to",
                "docx",
                "--standalone",
                "--metadata",
                "lang=de-DE",
                "-o",
                str(folder / target),
            ],
            check=True,
        )
        printable.unlink()
        order = list(range(61, 69)) + (
            list(range(61)) if source == "MANUSCRIPT.md" else []
        )
        chapter_titles = [
            next(
                line[2:]
                for line in (folder / f"section-{i:03d}.md")
                .read_text(encoding="utf-8")
                .splitlines()
                if line.startswith("# ")
            )
            for i in order
        ]
        style_docx(folder / target, title, chapter_titles)


def build(
    campaign: Path,
    folder: Path,
    *,
    docx: bool = False,
    pdf: bool = False,
    integrate: bool = False,
) -> None:
    campaign = campaign.resolve()
    plan, summary = load(campaign / "plan.json"), load(campaign / "summary.json")
    if not load(campaign / "completion.json")["source_unchanged"]:
        raise ValueError("Source changed during the campaign")
    for index, spec in enumerate(plan["selections"], 1):
        verified_result(campaign / f"{index:03d}-{spec['protocol']}", spec)
    if summary["accepted_evidence"] is not False:
        raise ValueError("Unexpected evidence authority")
    if folder.exists():
        raise ValueError("Refusing to overwrite an existing edition")
    previous = ROOT / "research/publications" / PREVIOUS
    folder.mkdir(parents=True)
    for path in sorted(previous.glob("section-*.md")):
        shutil.copy2(path, folder / path.name)
    for name in (
        "CITATION.bib",
        "CITATION.cff",
        "literatur_revision.bib",
        "methodenpruefung.py",
        "methodenpruefung.json",
    ):
        shutil.copy2(previous / name, folder / name)
    header = decode_text(
        r"""# Rekursive Epistemik in verk\u00f6rperten spikenden neuronalen Architekturen

## Vollst\u00e4ndiges Dissertationsmanuskript und wissenschaftliche Forschungsarbeit

**Thomas Heisig | Fassung 1.5 | 13. September 2026**

Technisches Framework: MHRN, Software-Entwicklungslinie 0.6.0-alpha.3. KI-unterst\u00fctzte Quellenpr\u00fcfung, Dokumentation und Interpretation. Menschliches wissenschaftliches Review: ausstehend. Keine automatische EVID-Freigabe und keine Behauptung akademischer Annahme.

## Editions- und Lesekontrakt

Diese Fassung enth\u00e4lt die vollst\u00e4ndige bisherige Abhandlung, nicht nur eine Zusammenfassung. Teil I aktualisiert Architektur, Methodik, tats\u00e4chliche Ausf\u00fchrung, Ergebnisse und Grenzen. Teil II bewahrt die 61 datierten Kapitel der Vorfassung als vollst\u00e4ndige theoretische und wissenschaftshistorische Grundlage. Zahlen und Statusbehauptungen aus den \u00e4lteren Kapiteln gelten f\u00fcr deren jeweilige Zeitpunkte; f\u00fcr den aktuellen Zustand hat Teil I Vorrang. Die Originalausgaben und Originaldaten werden nicht ver\u00e4ndert.

Die separate Forschungsarbeit in FORSCHUNGSBERICHT.md stellt den aktuellen Methoden-/Ergebnisteil eigenst\u00e4ndig bereit. Quellcommit, Digest, Protokolle, Seeds und Umgebung sind in der zugeordneten Kampagne dokumentiert. Vollst\u00e4ndigkeit einer Textausgabe bedeutet nicht, dass jede wissenschaftliche Hypothese bereits entschieden ist.

"""
    )
    software = f"# Quellenbindung und Zusammenfassung\n\nKampagne `{summary['campaign']}` auf Commit `{summary['source']}`.\n\n{len(plan['selections'])} ausgewiesene Ausfuehrungen; Status `{summary['status_counts']}`; {summary['total_runs']} gespeicherte Datensaetze. {len(plan['human_templates'])} menschliche Vorlagen werden nicht als Experimente ausgegeben. Die Unterscheidung zwischen Simulation, Grenzaudit, Komponentenfunktion und menschlicher Entscheidung bestimmt die Reichweite aller Aussagen.\n\nEin protokollierter Lauf, eine bestandene Softwarepruefung und eine wissenschaftliche Annahmeentscheidung bleiben verschiedene Objekte. Die Ausgabe dokumentiert auch verfehlte wissenschaftliche Erfolgskriterien.\n\n"
    chapters = {
        61: software,
        62: decode_text(ARCHITECTURE),
        63: decode_text(METHODS),
        64: results_text(plan, summary),
        65: decode_text(DISCUSSION),
        66: coverage_text(plan),
        67: decode_text(REVIEW),
        68: decode_text(REFERENCES),
    }
    for index, text in chapters.items():
        (folder / f"section-{index:03d}.md").write_text(
            text.strip() + "\n", encoding="utf-8"
        )
    order = [f"section-{i:03d}.md" for i in range(61, 69)] + [
        f"section-{i:03d}.md" for i in range(61)
    ]
    current = "\n\n".join(chapters.values())
    historical = "\n\n".join(
        (folder / name).read_text(encoding="utf-8") for name in order[8:]
    )
    (folder / "MANUSCRIPT.md").write_text(
        header
        + "\n# Teil I - Aktueller wissenschaftlicher Stand\n\n"
        + current
        + "\n\n# Teil II - Vollstaendige theoretische Grundlage und datierte Vorfassungen\n\nDie nachfolgenden Kapitel sind unveraenderte historische Uebernahmen. Fuer aktuelle Messergebnisse und offene Voraussetzungen ist Teil I massgeblich.\n\n"
        + historical,
        encoding="utf-8",
    )
    report_header = "# MHRN - Forschungsarbeit, Fassung 1.5\n\n**Thomas Heisig | 13. September 2026**\n\nEigenstaendige wissenschaftliche Forschungsarbeit mit Quellenstand, Architektur, Methoden, vollstaendiger aktueller Messbilanz, Diskussion und Grenzen. Menschliches Review ausstehend; keine automatische Evidenzfreigabe. Die historische Gesamtabhandlung ist separat im Dissertationsmanuskript enthalten.\n\n"
    (folder / "FORSCHUNGSBERICHT.md").write_text(
        report_header + current, encoding="utf-8"
    )
    # Full manuscript navigation remains separate from the standalone report.
    (folder / "README.md").write_text(
        header
        + "\n[Gesamtes Manuskript](MANUSCRIPT.md) | [Forschungsarbeit](FORSCHUNGSBERICHT.md) | [Manifest](manifest.json)\n\n"
        + (
            "[Dissertation (Word)](MHRN_Dissertationsmanuskript_v1.5.docx) | [Forschungsarbeit (Word)](MHRN_Forschungsbericht_v1.5.docx)\n\n"
            if docx or pdf
            else ""
        )
        + (
            "[Dissertation (PDF)](MHRN_Dissertationsmanuskript_v1.5.pdf) | [Forschungsarbeit (PDF)](MHRN_Forschungsbericht_v1.5.pdf)\n\n"
            if pdf
            else ""
        )
        + "\n".join(
            f"- [{next((line[2:] for line in (folder / name).read_text(encoding='utf-8').splitlines() if line.startswith('# ')), name)}]({name})"
            for name in order
        )
        + "\n",
        encoding="utf-8",
    )
    for name in ("CITATION.bib", "CITATION.cff"):
        path = folder / name
        path.write_text(
            path.read_text()
            .replace("1.4", "1.5")
            .replace("2026-09-10", "2026-09-13")
            .replace("Thomas-Heisig/Brain-5D", "Thomas-Heisig/MHRN"),
            encoding="utf-8",
        )
    if docx or pdf:
        export_docx(folder)
    if pdf:
        for document in sorted(folder.glob("*.docx")):
            with tempfile.TemporaryDirectory(prefix="mhrn-office-") as profile:
                subprocess.run(
                    [
                        "soffice",
                        "-env:UserInstallation=" + Path(profile).as_uri(),
                        "--headless",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        str(folder),
                        str(document),
                    ],
                    check=True,
                    timeout=180,
                )
            result = document.with_suffix(".pdf")
            if not result.is_file() or result.stat().st_size < 10000:
                raise ValueError(
                    "PDF export missing or unexpectedly empty: " + str(result)
                )
    manifest = {
        "schema_version": "1.0",
        "edition": "1.5",
        "date": "2026-09-13",
        "section_count": len(order),
        "section_order": order,
        "current_sections": order[:8],
        "historical_sections": order[8:],
        "inherited_edition": PREVIOUS,
        "inherited_files_sha256": files(previous),
        "inherited_manifest_sha256": sha(previous / "manifest.json"),
        "campaign": summary["campaign"],
        "source_commit": summary["source"],
        "campaign_summary_sha256": sha(campaign / "summary.json"),
        "campaign_plan_sha256": sha(campaign / "plan.json"),
        "authority": "interpretation_only",
        "human_review": "pending",
        "accepted_evidence": False,
        "automatic_evidence_promotion": False,
        "files_sha256": files(folder),
    }
    save(folder / "manifest.json", manifest)
    if integrate:
        if folder.resolve() != (ROOT / "research/publications" / EDITION).resolve():
            raise ValueError("Integrate requires canonical edition location")
        catalog_path = folder.parent / "catalog.json"
        catalog = load(catalog_path)
        for item in catalog["publications"]:
            item.update(current=False, edition_status="historical")
        identity = load(ROOT / "project_identity.json")
        pubid = "PUB-RECURSIVE-EPISTEMICS-20260913-V1.5"
        entry = f"publications/{EDITION}/README.md"
        catalog["publications"].insert(
            0,
            {
                "id": pubid,
                "title": identity["publication"]["title_en"],
                "author": "Thomas Heisig",
                "type": "wissenschaftliche_abhandlung",
                "version": "1.5",
                "date": "2026-09-13",
                "current": True,
                "edition_status": "current",
                "entrypoint": entry,
                "reader": entry,
                "snapshot": f"publications/{EDITION}",
                "manifest": f"publications/{EDITION}/manifest.json",
                "authority": "interpretation_only",
                "read_only": True,
                "automatic_evidence_promotion": False,
            },
        )
        catalog.update(
            current_publication_id=pubid,
            current_publication=pubid,
            current_entrypoint=entry,
        )
        save(catalog_path, catalog)
        identity["publication"].update(
            edition="1.5", path=str(folder.relative_to(ROOT))
        )
        save(ROOT / "project_identity.json", identity)
        index = folder.parent / "README.md"
        index.write_text(
            f"# Rekursive Epistemik - aktuelle Fassung 1.5\n\n[Dissertationsmanuskript und Forschungsarbeit]({EDITION}/README.md). Vollstaendige Ausgabe, aktuelle quellgebundene Messbilanz, menschliches Review ausstehend.\n\n<details><summary>Historischer Publikationsindex</summary>\n\n"
            + index.read_text()
            + "\n</details>\n",
            encoding="utf-8",
        )
        state = ROOT / "docs/05-quality/SCIENTIFIC_STATE_ALPHA3.md"
        state.write_text(
            software
            + decode_text(DISCUSSION)
            + "\n\nVollstaendige neue Fassung: ../../research/publications/"
            + EDITION
            + "/README.md\n",
            encoding="utf-8",
        )
        (ROOT / "docs/02-architecture/SCIENTIFIC_CONTRACTS_ALPHA3.md").write_text(
            decode_text(ARCHITECTURE), encoding="utf-8"
        )
        save(
            ROOT / "research/generated/EXPERIMENT_COVERAGE_ALPHA3.json",
            {
                "source_commit": summary["source"],
                "campaign": summary["campaign"],
                "selections": plan["selections"],
                "human_templates": plan["human_templates"],
                "accepted_evidence": False,
            },
        )
        (ROOT / "research/generated/SCIENTIFIC_AUDIT_ALPHA3.md").write_text(
            software + coverage_text(plan), encoding="utf-8"
        )
        for name in ("README.md", "docs/README.md", "research/README.md"):
            path = ROOT / name
            path.write_text(
                "<!-- alpha3-current-state -->\n## Wissenschaftlicher Stand: Alpha.3 / Publikationsfassung 1.5\n\n"
                + software.replace("# Quellenbindung und Zusammenfassung\n\n", "")
                + "Aktueller Einstieg: `research/publications/"
                + EDITION
                + "/README.md` (repository-relative). Detaillierte Architektur: `docs/02-architecture/SCIENTIFIC_CONTRACTS_ALPHA3.md`. Die nachfolgenden aelteren Statusabschnitte sind datierte Historie, keine aktuelle CI- oder EVID-Freigabe.\n\n"
                + path.read_text(),
                encoding="utf-8",
            )
        todo = ROOT / "docs/08-roadmap/TODO.md"
        todo.write_text(
            "## Alpha.3 - wissenschaftlicher Abgleich, 2026-09-13\n\n- [x] Die registrierte maschinelle Kampagne ausfuehren und alle Receipts/Negativbefunde bewahren.\n- [x] Vollstaendige Publikationsfassung 1.5 und eigenstaendige Forschungsarbeit erzeugen.\n- [x] Ausfuehrung, Grenzaudit und EVID-Freigabe getrennt dokumentieren.\n- [ ] Die "
            + str(len(plan["human_templates"]))
            + " menschlichen Vorlagen sowie die wissenschaftlichen Ergebnisreviews bearbeiten.\n- [ ] Direkte Messvertraege fuer die ausgewiesenen Grenzfragen und gekoppelte Kognitions-/Langzeitplastizitaetspruefungen vervollstaendigen.\n- [ ] Unabhaengige Replikation und konkrete EVID-Freigaben einholen.\n\nAeltere offene Zaehler unten sind zeitgebundene Bestandsaufnahmen; der aktuelle Registerabgleich steht in research/generated/EXPERIMENT_COVERAGE_ALPHA3.json.\n\n"
            + todo.read_text(),
            encoding="utf-8",
        )


def verify(folder: Path) -> None:
    manifest = load(folder / "manifest.json")
    if manifest["accepted_evidence"] or manifest["automatic_evidence_promotion"]:
        raise ValueError("Unreviewed edition cannot promote EVID")
    if files(folder) != manifest["files_sha256"]:
        raise ValueError("Edition file inventory or digest mismatch")
    previous = ROOT / "research/publications" / manifest["inherited_edition"]
    if sha(previous / "manifest.json") != manifest["inherited_manifest_sha256"]:
        raise ValueError("Historical publication manifest changed")
    if files(previous) != manifest["inherited_files_sha256"]:
        raise ValueError("Historical edition changed")
    for name in manifest["historical_sections"]:
        if sha(folder / name) != sha(previous / name):
            raise ValueError("Historical chapter not completely retained")
    campaign = ROOT / "research/experiments" / manifest["campaign"]
    if (
        sha(campaign / "summary.json") != manifest["campaign_summary_sha256"]
        or sha(campaign / "plan.json") != manifest["campaign_plan_sha256"]
    ):
        raise ValueError("Campaign provenance changed")
    for index, spec in enumerate(load(campaign / "plan.json")["selections"], 1):
        verified_result(campaign / f"{index:03d}-{spec['protocol']}", spec)
    print(
        "Verified complete edition",
        manifest["edition"],
        "sections",
        manifest["section_count"],
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--campaign", type=Path)
    parser.add_argument(
        "--output", type=Path, default=ROOT / "research/publications" / EDITION
    )
    parser.add_argument("--docx", action="store_true")
    parser.add_argument("--pdf", action="store_true")
    parser.add_argument("--integrate", action="store_true")
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    if args.verify:
        verify(args.output)
    elif args.campaign:
        build(
            args.campaign,
            args.output,
            docx=args.docx,
            pdf=args.pdf,
            integrate=args.integrate,
        )
    else:
        parser.error("Use --campaign or --verify")


if __name__ == "__main__":
    main()
