"""The new complete edition remains traceable DATA interpretation, never EVID."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from scripts.publication_alpha3 import EDITION, verify

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "research/publications" / EDITION


def test_alpha3_complete_edition_and_raw_receipts() -> None:
    verify(FOLDER)
    manifest = json.loads((FOLDER / "manifest.json").read_text())
    assert manifest["section_count"] == 69
    assert len(manifest["historical_sections"]) == 61
    assert len(manifest["current_sections"]) == 8
    assert manifest["human_review"] == "pending"
    assert manifest["authority"] == "interpretation_only"
    campaign = ROOT / "research/experiments" / manifest["campaign"]
    summary = json.loads((campaign / "summary.json").read_text())
    assert summary["status_counts"] == {"completed": 70}
    assert summary["total_runs"] == 2043
    assert summary["human_templates_pending"] == 25
    assert summary["accepted_evidence"] is False
    assert len(summary["protocols"]) == 70
    assert "brian2_single_neuron_v1" in (FOLDER / "FORSCHUNGSBERICHT.md").read_text()


def test_alpha3_office_pdf_and_populated_navigation() -> None:
    for name, minimum_sections in (
        ("MHRN_Dissertationsmanuskript_v1.5", 69),
        ("MHRN_Forschungsbericht_v1.5", 8),
    ):
        document = Document(FOLDER / (name + ".docx"))
        assert (
            len(
                document.element.xpath(
                    ".//w:bookmarkStart[starts-with(@w:name, 'mhrn_chapter_')]"
                )
            )
            == minimum_sections
        )
        assert "Inhaltsverzeichnis" in [p.text for p in document.paragraphs]
        assert not any(
            p.text.lstrip().startswith("# ")
            for p in document.paragraphs
            if p.style.style_id in ("BodyText", "FirstParagraph")
        )
        pdf = PdfReader(FOLDER / (name + ".pdf"))
        assert len(pdf.pages) > 10
        assert "Fassung 1.5" in pdf.pages[0].extract_text()
        assert "Inhaltsverzeichnis" in pdf.pages[1].extract_text()
    assert (FOLDER / "methodenpruefung.py").is_file()
    assert (FOLDER / "methodenpruefung.json").is_file()
